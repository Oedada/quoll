"""Файлы отчёта из готовых строк (reports-design §10).

без БД и без цикла событий: выполняется в пуле процессов рендера. Вход -
неизменяемые dataclass-ы, они переезжают в другой процесс
"""

import io
import math
import re
from dataclasses import dataclass
from datetime import date
from functools import partial
from pathlib import Path

import openpyxl
import xlwt
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    LongTable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from quoll.reports import labels
from quoll.reports.policy import Row, StatusKind

# шапка: заголовок, период, 7 фильтров, автор, дата, число строк, пустая строка.
# от неё зависит предел xls: 65 536 - шапка - заголовки - итог
XLS_HEADER_ROWS = 13
XLS_SHEET_ROWS = 65_536

_FONTS = Path(__file__).parent / "fonts"
pdfmetrics.registerFont(TTFont("DejaVu", _FONTS / "DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("DejaVu-Bold", _FONTS / "DejaVuSans-Bold.ttf"))

# ширины колонок Excel в символах
_XL_WIDTHS = {
    "university": 32,
    "direction": 16,
    "program": 24,
    "product": 22,
    "status": 26,
    "responsible": 24,
    "students": 11,
    "streams": 10,
    "teachers_kam": 14,
    "teachers_lms": 14,
    "transitions": 44,
    "contract_number": 16,
    "license_until": 12,
    "transfer_status": 13,
    "region": 20,
}
# по центру: числа, даты, короткие значения
_CENTERED = {
    "students",
    "streams",
    "teachers_kam",
    "teachers_lms",
    "license_until",
    "transfer_status",
}
_NUMERIC = {"students", "streams", "teachers_kam", "teachers_lms", "license_until"}

# палитра общая для всех форматов
_INK = "1F2937"
_MUTED = "6B7280"
_HEAD = "2F3B52"
_ZEBRA = "F4F6F9"
_LINE = "D5DAE1"
_TOTAL = "E6EAF0"
_STATUS_INK = {
    StatusKind.REFUSED: "B42318",
    StatusKind.DONE: "067647",
    StatusKind.AWAITING: "B54708",
}


@dataclass(frozen=True)
class Header:
    period: str
    filters: tuple[tuple[str, str], ...]
    author: str
    formed_at: str
    row_count: int

    def about(self) -> list[tuple[str, str]]:
        return [
            ("Период", self.period),
            ("Сформирован", self.formed_at),
            ("Автор", self.author),
            ("Строк", str(self.row_count)),
        ]

    def items(self) -> list[tuple[str, str]]:
        """ровно XLS_HEADER_ROWS - 2 строк: всё, кроме заголовка и пустой"""
        return [*self.about()[:1], *self.filters, *self.about()[1:]]


def cells(row: Row, columns: list[str]) -> list:
    """значения колонок: числа и даты - как есть, пустое - прочерк"""
    values = {
        "university": row.university,
        "direction": row.direction or labels.NOT_SET["direction"],
        "program": row.program or labels.NOT_SET["program"],
        "product": row.product or labels.NOT_SET["product"],
        "status": row.status.label,
        "responsible": labels.responsible(row.responsible_name, row.earlier_name),
        "students": None,
        "streams": None,
        "teachers_kam": row.teachers_kam,
        "teachers_lms": None,
        "transitions": labels.moves(row.moves, "\n"),
        "contract_number": row.contract_number,
        "license_until": row.license_until,
        "transfer_status": labels.TRANSFER.get(row.transfer_status),
        "region": row.region,
    }
    return [labels.DASH if values[c] is None else values[c] for c in columns]


def _total(rows: list[Row], columns: list[str]) -> list | None:
    """итог по числу обученных КАМом; LMS не суммируются - данных нет (В1)"""
    if "teachers_kam" not in columns:
        return None
    total = sum(r.teachers_kam or 0 for r in rows)
    return ["Итого", *(total if c == "teachers_kam" else "" for c in columns[1:])]


def _row_height(values: list, widths: list[int], line: float) -> float:
    """высота строки Excel по переносам: сам Excel её при открытии не считает"""
    lines = 1
    for value, width in zip(values, widths, strict=True):
        if isinstance(value, str):
            # Calibri и Arial 10 уже цифры шрифта по умолчанию, в ширине их больше
            per_line = max(1, int(width * 1.15) - 1)
            lines = max(
                lines,
                sum(max(1, math.ceil(len(p) / per_line)) for p in value.split("\n")),
            )
    return min(409, lines * line + 4)


# --- xlsx


class _XlsxStyles:
    """стили создаются один раз: на каждую ячейку - только ссылки"""

    def __init__(self):
        thin = Side(style="thin", color=_LINE)
        self.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        self.total_border = Border(
            left=thin, right=thin, top=Side(style="medium", color=_HEAD), bottom=thin
        )
        self.font = Font(name="Calibri", size=10, color=_INK)
        self.muted = Font(name="Calibri", size=10, color=_MUTED)
        self.bold = Font(name="Calibri", size=10, bold=True, color=_INK)
        self.title = Font(name="Calibri", size=14, bold=True, color=_HEAD)
        self.label = Font(name="Calibri", size=10, bold=True, color=_MUTED)
        self.head = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        self.status = {
            kind: Font(name="Calibri", size=10, bold=True, color=ink)
            for kind, ink in _STATUS_INK.items()
        }
        self.head_fill = PatternFill("solid", fgColor=_HEAD)
        self.zebra = PatternFill("solid", fgColor=_ZEBRA)
        self.total_fill = PatternFill("solid", fgColor=_TOTAL)
        self.left = Alignment(horizontal="left", vertical="top", wrap_text=True)
        self.center = Alignment(horizontal="center", vertical="top", wrap_text=True)
        self.head_align = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )


def xlsx(header: Header, rows: list[Row], columns: list[str]) -> bytes:
    book = openpyxl.Workbook(write_only=True)
    sheet = book.create_sheet("Отчёт")
    st = _XlsxStyles()
    widths = [_XL_WIDTHS[c] for c in columns]
    for n, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(n)].width = width
    top = XLS_HEADER_ROWS + 1
    last = get_column_letter(len(columns))
    # закреплена строка заголовков - первая строка под ней
    sheet.freeze_panes = f"A{top + 1}"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = f"A{top}:{last}{top + len(rows)}"
    sheet.print_title_rows = f"{top}:{top}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = 9  # A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.oddFooter.left.text = labels.TITLE
    sheet.oddFooter.right.text = "стр. &P из &N"

    title = WriteOnlyCell(sheet, value=labels.TITLE)
    title.font = st.title
    sheet.row_dimensions[1].height = 24
    sheet.append([title])
    for name, value in header.items():
        label = WriteOnlyCell(sheet, value=name)
        label.font = st.label
        text = WriteOnlyCell(sheet, value=value)
        text.font = st.font
        sheet.append([label, text])
    sheet.append([])

    titles = []
    for c in columns:
        cell = WriteOnlyCell(sheet, value=labels.COLUMNS[c])
        cell.font, cell.fill = st.head, st.head_fill
        cell.alignment, cell.border = st.head_align, st.border
        titles.append(cell)
    sheet.row_dimensions[top].height = _row_height(
        [labels.COLUMNS[c] for c in columns], widths, 13
    )
    sheet.append(titles)

    for n, row in enumerate(rows):
        values = cells(row, columns)
        out = []
        for c, value in zip(columns, values, strict=True):
            cell = WriteOnlyCell(sheet, value=value)
            cell.border = st.border
            cell.alignment = st.center if c in _CENTERED else st.left
            if value == labels.DASH:
                cell.font, cell.alignment = st.muted, st.center
            elif c == "status" and row.status.kind in st.status:
                cell.font = st.status[row.status.kind]
            else:
                cell.font = st.font
            if isinstance(value, date):
                cell.number_format = "DD.MM.YYYY"
            if n % 2:
                cell.fill = st.zebra
            out.append(cell)
        sheet.row_dimensions[top + 1 + n].height = _row_height(values, widths, 13)
        sheet.append(out)

    if (total := _total(rows, columns)) is not None:
        out = []
        for c, value in zip(columns, total, strict=True):
            cell = WriteOnlyCell(sheet, value=value)
            cell.font, cell.fill, cell.border = st.bold, st.total_fill, st.total_border
            cell.alignment = st.center if c in _CENTERED else st.left
            out.append(cell)
        sheet.append(out)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# --- xls


class _XlsStyles:
    """xlwt: палитра на 56 цветов и лимит стилей - всё создаётся один раз"""

    def __init__(self, book: xlwt.Workbook):
        for n, (name, rgb) in enumerate(
            [
                ("report_ink", _INK),
                ("report_muted", _MUTED),
                ("report_head", _HEAD),
                ("report_zebra", _ZEBRA),
                ("report_line", _LINE),
                ("report_total", _TOTAL),
                *((f"report_{k.value.lower()}", v) for k, v in _STATUS_INK.items()),
            ],
            start=0x20,
        ):
            xlwt.add_palette_colour(name, n)
            book.set_colour_RGB(n, *bytes.fromhex(rgb))
        font = "font: name Arial, height 200"
        borders = (
            "borders: left thin, right thin, top thin, bottom thin, "
            "left_colour report_line, right_colour report_line, "
            "top_colour report_line, bottom_colour report_line"
        )
        self.title = xlwt.easyxf(
            "font: name Arial, height 280, bold on, colour report_head"
        )
        self.label = xlwt.easyxf(f"{font}, bold on, colour report_muted")
        self.value = xlwt.easyxf(f"{font}, colour report_ink")
        self.head = xlwt.easyxf(
            f"{font}, bold on, colour white; {borders}; "
            "pattern: pattern solid, fore_colour report_head; "
            "align: wrap on, vert centre, horiz centre"
        )
        self.cells = {}
        for zebra in (False, True):
            fill = "; pattern: pattern solid, fore_colour report_zebra" if zebra else ""
            for kind in ("left", "center", "muted", "date", *_STATUS_INK):
                colour = {
                    "muted": "report_muted",
                    **{k: f"report_{k.value.lower()}" for k in _STATUS_INK},
                }.get(kind, "report_ink")
                bold = ", bold on" if kind in _STATUS_INK else ""
                horiz = "centre" if kind in ("center", "muted", "date") else "left"
                self.cells[kind, zebra] = xlwt.easyxf(
                    f"{font}, colour {colour}{bold}; {borders}; "
                    f"align: wrap on, vert top, horiz {horiz}{fill}",
                    num_format_str="DD.MM.YYYY" if kind == "date" else "General",
                )
        self.total = {
            horiz: xlwt.easyxf(
                f"{font}, bold on, colour report_ink; {borders}, "
                "top medium, top_colour report_head; "
                "pattern: pattern solid, fore_colour report_total; "
                f"align: wrap on, vert top, horiz {horiz}"
            )
            for horiz in ("left", "centre")
        }


def xls(header: Header, rows: list[Row], columns: list[str]) -> bytes:
    book = xlwt.Workbook(encoding="utf-8")
    sheet = book.add_sheet("Отчёт")
    st = _XlsStyles(book)
    widths = [_XL_WIDTHS[c] for c in columns]
    for n, width in enumerate(widths):
        sheet.col(n).width = 256 * width
    sheet.show_grid = False
    sheet.portrait = False
    sheet.paper_size_code = 9  # A4
    sheet.fit_num_pages = 0
    sheet.fit_width_to_pages = 1
    sheet.fit_height_to_pages = 0
    # xlwt на python 3 пишет в колонтитул только ASCII
    sheet.header_str = b""
    sheet.footer_str = b"&R&P / &N"

    def height(n: int, points: float) -> None:
        sheet.row(n).height_mismatch = True
        sheet.row(n).height = int(points * 20)

    sheet.write(0, 0, labels.TITLE, st.title)
    height(0, 24)
    for n, (name, value) in enumerate(header.items(), start=1):
        sheet.write(n, 0, name, st.label)
        sheet.write(n, 1, value, st.value)
    top = XLS_HEADER_ROWS
    for n, c in enumerate(columns):
        sheet.write(top, n, labels.COLUMNS[c], st.head)
    height(top, _row_height([labels.COLUMNS[c] for c in columns], widths, 13))
    sheet.set_panes_frozen(True)
    sheet.set_horz_split_pos(top + 1)
    for r, row in enumerate(rows, start=top + 1):
        values = cells(row, columns)
        zebra = bool((r - top - 1) % 2)
        for n, (c, value) in enumerate(zip(columns, values, strict=True)):
            if value == labels.DASH:
                kind = "muted"
            elif isinstance(value, date):
                kind = "date"
            elif c == "status" and row.status.kind in _STATUS_INK:
                kind = row.status.kind
            else:
                kind = "center" if c in _CENTERED else "left"
            sheet.write(r, n, value, st.cells[kind, zebra])
        height(r, _row_height(values, widths, 13))
    if (total := _total(rows, columns)) is not None:
        r = top + 1 + len(rows)
        for n, (c, value) in enumerate(zip(columns, total, strict=True)):
            sheet.write(r, n, value, st.total["centre" if c in _CENTERED else "left"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# --- pdf

_PAGE = landscape(A4)
_MARGIN = 12 * mm
# кегль таблицы: крупнейший, при котором слова влезают в колонки
_BODY_SIZES = (7.5, 7, 6.5, 6)
_PAD = 3
_HEAD_RATIO = 7 / 7.5
# предел ширины колонки pdf: длинный текст переносится, а не растягивает её
_PDF_CAPS = {"transitions": 70 * mm, "university": 42 * mm, "status": 44 * mm}
_PDF_CAP = 40 * mm
# желаемый минимум: строка до этой ширины не переносится, пока хватает места.
# Переход, статус и продукт в одну строку читаются заметно легче
_PDF_SOFT = {
    "transitions": 52 * mm,
    "status": 34 * mm,
    "product": 28 * mm,
    "university": 28 * mm,
    "program": 24 * mm,
}
# при нехватке места эти колонки ужимаются последними: перенос внутри
# перехода или статуса читается хуже, чем в названии вуза
_PDF_SHRINK = {"transitions": 0.3, "status": 0.6, "responsible": 0.8}


class _NumberedCanvas(canvas.Canvas):
    """колонтитул и "стр. N из M": страницы копятся, номера - когда M известно"""

    def __init__(self, *args, footer: str = "", **kwargs):
        super().__init__(*args, **kwargs)
        self._footer = footer
        self._pages = []

    def showPage(self):  # noqa: N802 - имя из reportlab
        self._pages.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        count = len(self._pages)
        width = self._pagesize[0]
        for state in self._pages:
            self.__dict__.update(state)
            self.setStrokeColor(colors.HexColor(f"#{_LINE}"))
            self.setLineWidth(0.5)
            self.line(_MARGIN, 9 * mm, width - _MARGIN, 9 * mm)
            self.setFillColor(colors.HexColor(f"#{_MUTED}"))
            self.setFont("DejaVu", 7)
            self.drawString(_MARGIN, 6 * mm, self._footer)
            self.drawRightString(
                width - _MARGIN, 6 * mm, f"стр. {self._pageNumber} из {count}"
            )
            super().showPage()
        super().save()


def _styles(size: float) -> dict[str, ParagraphStyle]:
    body = ParagraphStyle(
        "body",
        fontName="DejaVu",
        fontSize=size,
        leading=size * 1.3,
        textColor=colors.HexColor(f"#{_INK}"),
    )
    bold = ParagraphStyle("bold", parent=body, fontName="DejaVu-Bold")
    return {
        "body": body,
        "center": ParagraphStyle("center", parent=body, alignment=TA_CENTER),
        "bold": bold,
        "bold_center": ParagraphStyle("bold_center", parent=bold, alignment=TA_CENTER),
        "head": ParagraphStyle(
            "head",
            parent=bold,
            alignment=TA_CENTER,
            textColor=colors.white,
            fontSize=size - 0.5,
            leading=(size - 0.5) * 1.25,
        ),
        "title": ParagraphStyle(
            "title",
            parent=bold,
            fontSize=15,
            leading=19,
            textColor=colors.HexColor(f"#{_HEAD}"),
        ),
        "label": ParagraphStyle(
            "label",
            parent=bold,
            fontSize=7.5,
            leading=9.75,
            textColor=colors.HexColor(f"#{_MUTED}"),
        ),
        "value": ParagraphStyle("value", parent=body, fontSize=7.5, leading=9.75),
    }


_NBSP = "\u00a0"
# склеиваются только короткие пары - длинное неразрывное слово растянуло бы колонку
_GLUE_MAX = 12


def _short(m: re.Match) -> str:
    text = m.group(0)
    return text.replace(" ", _NBSP) if len(text) <= _GLUE_MAX + 3 else text


# неразрывные пары: номер шага с названием, короткое в скобках, фамилия
# с инициалами, инициалы, "не ..."
_GLUE = [
    (re.compile(r"\b(\d+(?:\.\d+)?\.?) (?=[А-ЯЁA-Z])"), rf"\1{_NBSP}"),
    (re.compile(r"\([^()]*\)"), _short),
    (re.compile(r"\b[^\W\d_]+ (?=[^\W\d_]\.(?:\s|$))"), _short),
    (re.compile(r"(\b[^\W\d_]\.) (?=[^\W\d_]\.)"), rf"\1{_NBSP}"),
    (re.compile(r"\bне "), f"не{_NBSP}"),
]


def _glue(text: str) -> str:
    for pattern, repl in _GLUE:
        text = pattern.sub(repl, text)
    return text


def _esc(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _muted(text: str) -> str:
    return f'<font color="#{_MUTED}">{_esc(text)}</font>'


def _pdf_text(value) -> str:
    return value.strftime("%d.%m.%Y") if isinstance(value, date) else str(value)


def _pdf_cell(row: Row, column: str, value) -> tuple[str, list[str]]:
    """разметка ячейки и её строки без разметки - по ним меряется ширина"""
    if value == labels.DASH:
        return _muted(labels.DASH), [labels.DASH]
    if isinstance(value, str):
        value = _glue(value)
    if column == "status":
        main, *extra = value.split(" · ")
        ink = _STATUS_INK.get(row.status.kind)
        text = f'<font color="#{ink}"><b>{_esc(main)}</b></font>' if ink else _esc(main)
        return text + "".join(_muted(f" · {x}") for x in extra), [value]
    if column == "responsible" and row.earlier_name:
        name = _glue(row.responsible_name)
        earlier = _glue(f"ранее: {row.earlier_name}")
        return f"{_esc(name)}<br/>{_muted(earlier)}", [name, earlier]
    if column == "transitions":
        if not row.moves:
            return _muted(labels.NO_MOVES), [labels.NO_MOVES]
        lines = [_glue(labels.move(m)) for m in row.moves]
        marked = []
        for line in lines:
            day, rest = line.split(" ", 1)
            marked.append(f"{_muted(day)} {_esc(rest)}")
        return "<br/>".join(marked), lines
    text = _pdf_text(value)
    return _esc(text), text.split("\n")


def _measure(columns: list[str], texts: list[list[list[str]]]):
    """самое длинное слово и самая длинная строка колонки при кегле 1:
    ширина текста пропорциональна кеглю, мерить второй раз не нужно"""
    measure = partial(pdfmetrics.stringWidth, fontSize=1)
    words, lines = [], []
    for n, c in enumerate(columns):
        head = labels.PDF_COLUMNS[c].split(" ")
        # заголовок на полкегля меньше и жирный
        low = max(measure(w, "DejaVu-Bold") for w in head) * _HEAD_RATIO
        high = low
        for cell in texts[n]:
            for line in cell:
                high = max(high, measure(line, "DejaVu"))
                low = max(low, max(measure(w, "DejaVu") for w in line.split(" ")))
        words.append(low)
        lines.append(high)
    return words, lines


def _widths(columns: list[str], words, lines, size: float, avail: float, soft: bool):
    """ширины по тексту: не уже самого длинного слова, не шире нужного.
    Лишнее место уходит текстовым колонкам, нехватка ужимает переносы.
    None - при этом кегле слова в страницу не влезают"""
    pad = 2 * _PAD + 1
    mins = [w * size + pad for w in words]
    if soft:
        mins = [
            max(lo, min(hi * size + pad, _PDF_SOFT.get(c, 0)))
            for c, lo, hi in zip(columns, mins, lines, strict=True)
        ]
    if sum(mins) > avail:
        return None
    prefs = [
        max(lo, min(hi * size + pad, _PDF_CAPS.get(c, _PDF_CAP)))
        for c, lo, hi in zip(columns, mins, lines, strict=True)
    ]
    if sum(prefs) <= avail:
        grow = [p if c not in _NUMERIC else 0 for c, p in zip(columns, prefs)]
        extra = avail - sum(prefs)
        return [p + extra * g / (sum(grow) or 1) for p, g in zip(prefs, grow)]
    widths = list(prefs)
    weights = [_PDF_SHRINK.get(c, 1.0) for c in columns]
    deficit = sum(widths) - avail
    while deficit > 0.01:
        slack = [w * (x - lo) for w, x, lo in zip(weights, widths, mins)]
        share = sum(slack)
        cut = [
            min(x - lo, deficit * sl / share) for x, lo, sl in zip(widths, mins, slack)
        ]
        widths = [x - c for x, c in zip(widths, cut)]
        deficit -= sum(cut)
    return widths


def _pdf_header(header: Header, st, width: float) -> Table:
    """сведения об отчёте слева, фильтры справа"""
    left, right = header.about(), list(header.filters)
    data = []
    for n in range(max(len(left), len(right))):
        line = []
        for part in (left, right):
            if n < len(part):
                name, value = part[n]
                line += [
                    Paragraph(_esc(name), st["label"]),
                    Paragraph(_esc(value), st["value"]),
                ]
            else:
                line += ["", ""]
        data.append(line)
    table = Table(
        data,
        colWidths=[27 * mm, width * 0.42 - 27 * mm, 30 * mm, width * 0.58 - 30 * mm],
    )
    table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                # пустые ячейки - того же кегля, иначе строки разной высоты
                ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                ("LEADING", (0, 0), (-1, -1), 9.75),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 1),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
            ]
        )
    )
    return table


def pdf(header: Header, rows: list[Row], columns: list[str]) -> bytes:
    buffer = io.BytesIO()
    avail = _PAGE[0] - 2 * _MARGIN
    doc = SimpleDocTemplate(
        buffer,
        pagesize=_PAGE,
        leftMargin=_MARGIN,
        rightMargin=_MARGIN,
        topMargin=_MARGIN,
        bottomMargin=14 * mm,
        title=labels.TITLE,
        author=header.author,
    )
    marked = [[] for _ in columns]
    plain = [[] for _ in columns]
    for row in rows:
        for n, (c, value) in enumerate(zip(columns, cells(row, columns), strict=True)):
            text, lines = _pdf_cell(row, c, value)
            marked[n].append(text)
            plain[n].append(lines)
    words, lines = _measure(columns, plain)
    # сначала крупный кегль с желаемыми минимумами, потом только по словам
    for soft, size in [(True, s) for s in _BODY_SIZES] + [
        (False, s) for s in _BODY_SIZES
    ]:
        widths = _widths(columns, words, lines, size, avail, soft)
        if widths is not None:
            break
    else:
        # слова не влезают и при мелком кегле - длинные рвутся по буквам
        mins = [w * size + 2 * _PAD + 1 for w in words]
        widths = [lo * avail / sum(mins) for lo in mins]
    st = _styles(size)

    data = [[Paragraph(_esc(labels.PDF_COLUMNS[c]), st["head"]) for c in columns]]
    for r in range(len(rows)):
        data.append(
            [
                Paragraph(marked[n][r], st["center" if c in _CENTERED else "body"])
                for n, c in enumerate(columns)
            ]
        )
    total = _total(rows, columns)
    if total is not None:
        data.append(
            [
                Paragraph(_esc(str(v)), st["bold_center" if c in _CENTERED else "bold"])
                for c, v in zip(columns, total, strict=True)
            ]
        )
    table = LongTable(data, colWidths=widths, repeatRows=1)
    line = colors.HexColor(f"#{_LINE}")
    commands = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("VALIGN", (0, 0), (-1, 0), "MIDDLE"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_HEAD}")),
        ("LEFTPADDING", (0, 0), (-1, -1), _PAD),
        ("RIGHTPADDING", (0, 0), (-1, -1), _PAD),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
        ("LINEBELOW", (0, 1), (-1, -1), 0.4, line),
        ("LINEAFTER", (0, 1), (-2, -1), 0.25, line),
        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor(f"#{_HEAD}")),
        (
            "ROWBACKGROUNDS",
            (0, 1),
            (-1, len(rows)),
            [colors.white, colors.HexColor(f"#{_ZEBRA}")],
        ),
    ]
    if total is not None:
        commands += [
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor(f"#{_TOTAL}")),
            ("LINEABOVE", (0, -1), (-1, -1), 0.9, colors.HexColor(f"#{_HEAD}")),
        ]
    table.setStyle(TableStyle(commands))

    story = [
        Paragraph(labels.TITLE, st["title"]),
        Spacer(1, 2.5 * mm),
        _pdf_header(header, st, avail),
        Spacer(1, 5 * mm),
        table,
    ]
    footer = f"{labels.TITLE} · {header.period} · сформирован {header.formed_at}"
    doc.build(story, canvasmaker=partial(_NumberedCanvas, footer=footer))
    return buffer.getvalue()


RENDERERS = {"xlsx": xlsx, "xls": xls, "pdf": pdf}
CONTENT_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "xls": "application/vnd.ms-excel",
    "pdf": "application/pdf",
}
